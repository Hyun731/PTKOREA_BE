"""Native PPTX renderer (PRD 11): text, tables and charts are editable PowerPoint objects,
images keep the full original with a non-destructive crop, every shape is named after its
block_id so QC findings point at an editable region, and notes carry evidence IDs.
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable, Optional

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from core.composition.tokens import color, resolve_weight, tokens, weight_info
from core.schemas.slidespec import (
    BarChartBlock, DeckSpec, ImageBlock, LineChartBlock, Paragraph, Region, ShapeBlock, SlideSpec, TableBlock, TextBlock,
)

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
ROLE_COLORS = {"self": "navy", "key": "red", "other": "bar_other", "other2": "line_other2"}
LINE_ROLE_COLORS = {"self": "navy", "key": "red", "other": "line_other", "other2": "line_other2"}

AssetResolver = Callable[[str], tuple[str, int, int]]  # asset_id -> (path, width, height)


def _emu(v: float) -> Emu:
    return Emu(int(round(v * tokens()["canvas"]["emu_per_unit"])))


def _rgb(role: str) -> RGBColor:
    return RGBColor.from_string(color(role))


def _font_family() -> str:
    return tokens()["font"]["family"]


def _set_run_font(run, size: float, weight, rgb: str, family: Optional[str] = None):
    """weight: a token weight name (light/regular/semibold/bold/heavy) or a legacy bool."""
    info = weight_info(weight)
    family = family or info["family"]
    font = run.font
    font.size = Pt(size)
    font.bold = info["bold"]
    font.color.rgb = RGBColor.from_string(rgb)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", family)
    rPr.set("lang", "ko-KR")
    rPr.set("altLang", "en-US")


def _fill_paragraphs(tf, paragraphs: list[Paragraph], *, size: float, line: float, align: str, color_role: str,
                     bold: bool, weight: Optional[str] = None):
    tf.clear()
    first = True
    for para in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[align]
        p.line_spacing = Pt(line)
        p.space_before = Pt(para.space_before)
        p.space_after = Pt(0)
        # Do not break Latin/number runs such as "−2.3pp" in the middle.
        p._p.get_or_add_pPr().set("latinLnBrk", "0")
        base_color = para.color or color_role
        if para.bullet:
            r = p.add_run()
            r.text = para.bullet + " "
            on_dark = base_color in ("white", "summary_text", "slate", "mist", "gold")
            bullet_color = ("red" if para.bullet == "▶" else
                            ("navy" if para.bullet in ("┃", "▲") and not on_dark else base_color))
            bw = "bold" if para.bullet == "┃" else resolve_weight(None, False, para.weight, para.bold, weight, bold)
            _set_run_font(r, size, bw, color(bullet_color))
        for part in para.parts:
            text = part.rendered()
            if not text:
                continue
            r = p.add_run()
            r.text = text
            _set_run_font(r, part.size or size, resolve_weight(part.weight, part.bold, para.weight, para.bold, weight, bold),
                          color(part.color or base_color))
            if part.type == "link" and part.href:
                r.hyperlink.address = part.href
                r.font.underline = True


def _text(slide, b: TextBlock):
    reg = b.region
    if b.role == "takeaway":
        # v6: the summary line sits on a navy rule that heads the numbered insights.
        rule_y = tokens()["regions"].get("takeaway_rule_y", reg.y + reg.h)
        rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _emu(reg.x), _emu(rule_y), _emu(reg.x + reg.w), _emu(rule_y))
        rule.name = f"{b.block_id}_rule"
        _strip_style(rule)
        rule.line.color.rgb = _rgb("navy")
        rule.line.width = Pt(1)
        box = slide.shapes.add_textbox(_emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h))
    elif b.fill:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h))
        _solid(box, b.fill)
    else:
        box = slide.shapes.add_textbox(_emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h))
    box.name = b.block_id
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[b.valign]
    _fill_paragraphs(tf, b.paragraphs, size=b.size, line=b.line_spacing, align=b.align, color_role=b.color, bold=b.bold,
                     weight=b.weight)


def _strip_style(shape):
    """Drop the theme style reference so no theme shadow/outline leaks into the output."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _solid(shape, role: Optional[str], line: Optional[str] = None, line_width: float = 0.75):
    _strip_style(shape)
    if role:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(role)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()


def _shape(slide, b: ShapeBlock):
    reg = b.region
    if b.shape == "line":
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _emu(reg.x), _emu(reg.y), _emu(reg.x + reg.w), _emu(reg.y + reg.h))
        _strip_style(ln)
        ln.line.color.rgb = _rgb(b.line or "muted")
        ln.line.width = Pt(b.line_width)
        ln.name = b.block_id
        return
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if b.shape == "rounded" else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, _emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h))
    if b.shape == "rounded":
        shp.adjustments[0] = 0.18
    shp.name = b.block_id
    _solid(shp, b.fill, b.line, b.line_width)


def _chart_text(chart, size: float, rgb: str = "595959"):
    chart.font.size = Pt(size)
    chart.font.name = _font_family()
    chart.font.color.rgb = RGBColor.from_string(rgb)


def _bar_chart(slide, b: BarChartBlock):
    reg = b.region
    data = CategoryChartData()
    data.categories = [it.label for it in b.items]
    data.add_series("점유(%)", [it.value for it in b.items], number_format=b.number_format)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, _emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h), data)
    frame.name = b.block_id
    chart = frame.chart
    _chart_text(chart, tokens()["type"]["chart_label"]["size"], color("subtle"))
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = b.gap_width
    plot.vary_by_categories = False
    cat = chart.category_axis
    cat.reverse_order = True  # first item on top
    cat.has_major_gridlines = False
    cat.format.line.fill.background()
    cat.tick_labels.font.size = Pt(b.label_size)
    cat.tick_labels.font.color.rgb = _rgb("text")
    val = chart.value_axis
    top = max((it.value for it in b.items), default=1.0)
    val.maximum_scale = top * 1.12 if top > 0 else 1
    val.minimum_scale = 0
    val.has_major_gridlines = False
    val.visible = False
    series = plot.series[0]
    # Series-level labels: a per-point override below would otherwise hide plot-level labels.
    labels = series.data_labels
    labels.show_value = True
    labels.number_format = b.number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.INSIDE_END
    labels.font.size = Pt(b.label_size + 0.5)
    labels.font.bold = True
    labels.font.color.rgb = _rgb("white")
    for i, it in enumerate(b.items):
        point = series.points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(ROLE_COLORS[it.role])
        point.format.line.fill.background()
        # Short bars cannot hold a white label: put it outside in dark text (reference p.26).
        if top > 0 and it.value / (top * 1.12) < 0.22:
            dl = point.data_label
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.size = Pt(b.label_size + 0.5)
            dl.font.bold = True
            dl.font.color.rgb = _rgb("text")
            _show_value(dl)


def _show_value(data_label):
    """python-pptx creates a bare dLbl; make it show the value with our number format."""
    dLbl = data_label._get_or_add_dLbl()
    for tag, val in (("c:showLegendKey", "0"), ("c:showVal", "1"), ("c:showCatName", "0"), ("c:showSerName", "0"),
                     ("c:showPercent", "0"), ("c:showBubbleSize", "0")):
        el = dLbl.find(qn(tag))
        if el is None:
            el = etree.SubElement(dLbl, qn(tag))
        el.set("val", val)


def _line_chart(slide, b: LineChartBlock):
    reg = b.region
    data = CategoryChartData()
    data.categories = b.categories
    for s in b.series:
        data.add_series(s.entity, [p.value for p in s.points], number_format="0.0")
    frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, _emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h), data)
    frame.name = b.block_id
    chart = frame.chart
    _chart_text(chart, 9, color("subtle"))
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9)
    val = chart.value_axis
    val.has_major_gridlines = False
    val.visible = False
    vals = [p.value for s in b.series for p in s.points if p.value is not None]
    if vals:
        val.minimum_scale = 0
        val.maximum_scale = max(vals) * 1.1
    cat = chart.category_axis
    cat.has_major_gridlines = False
    cat.format.line.color.rgb = _rgb("table_border")
    cat.tick_labels.font.size = Pt(9.5)
    # Missing months stay gaps (never interpolated).
    plot_area = chart._chartSpace.find(qn("c:chart"))
    disp = plot_area.find(qn("c:dispBlanksAs"))
    if disp is None:
        disp = etree.SubElement(plot_area, qn("c:dispBlanksAs"))
    disp.set("val", "gap")
    for series, spec in zip(chart.plots[0].series, b.series):
        rgb = _rgb(LINE_ROLE_COLORS[spec.role])
        series.smooth = False
        series.format.line.color.rgb = rgb
        series.format.line.width = Pt(2.5 if spec.role in ("self", "key") else 2)
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 6
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = rgb
        series.marker.format.line.color.rgb = rgb


def _cell_border(cell, rgb: str, width_pt: float = 0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for idx, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(int(width_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=rgb)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tcPr.insert(idx, ln)


def _table(slide, b: TableBlock):
    reg = b.region
    n_rows = len(b.rows) + 1
    n_cols = len(b.columns)
    frame = slide.shapes.add_table(n_rows, n_cols, _emu(reg.x), _emu(reg.y), _emu(reg.w), _emu(reg.h))
    frame.name = b.block_id
    tbl = frame.table
    # No built-in table style: every fill/border below is explicit and editable.
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    widths = [c.width for c in b.columns]
    total = sum(widths)
    acc = 0
    for i, c in enumerate(b.columns):
        w = reg.w * c.width / total
        tbl.columns[i].width = _emu(w)
        acc += w
    tbl.rows[0].height = _emu(b.header_height)
    for r in range(1, n_rows):
        tbl.rows[r].height = _emu(b.row_height)
    border = color("table_border")

    def setup(cell, fill: Optional[str]):
        cell.margin_left = cell.margin_right = _emu(6)
        cell.margin_top = cell.margin_bottom = _emu(2)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        _cell_border(cell, border)
        if fill:
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb("white")

    for ci, col in enumerate(b.columns):
        cell = tbl.cell(0, ci)
        setup(cell, "navy")
        _fill_paragraphs(cell.text_frame, [Paragraph(parts=[{"type": "text", "text": col.label}])], size=b.header_size,
                         line=b.header_size * 1.35, align="c", color_role="white", bold=False, weight="semibold")
    for ri, row in enumerate(b.rows, start=1):
        for ci, (col, cell_spec) in enumerate(zip(b.columns, row.cells)):
            cell = tbl.cell(ri, ci)
            if cell_spec.merged:
                continue
            setup(cell, cell_spec.fill)
            role = cell_spec.color or ("muted" if row.muted else "text")
            _fill_paragraphs(cell.text_frame, cell_spec.paragraphs or [Paragraph()], size=b.size, line=b.size * 1.4,
                             align=cell_spec.align or col.align, color_role=role, bold=cell_spec.bold)
            cell.text_frame.word_wrap = True
    for ri, row in enumerate(b.rows, start=1):
        for ci, cell_spec in enumerate(row.cells):
            if cell_spec.merge_down:
                tbl.cell(ri, ci).merge(tbl.cell(ri + cell_spec.merge_down, ci))


def _image(slide, b: ImageBlock, resolve: AssetResolver):
    path, width, height = resolve(b.asset_id)
    c = b.crop
    crop_w_px, crop_h_px = width * c.w, height * c.h
    reg = b.region
    # Contain: never crop beyond the person's selection (subject and in-image text stay).
    scale = min(reg.w / crop_w_px, reg.h / crop_h_px)
    w, h = crop_w_px * scale, crop_h_px * scale
    x, y = reg.x + (reg.w - w) / 2, reg.y + (reg.h - h) / 2
    pic = slide.shapes.add_picture(path, _emu(x), _emu(y), _emu(w), _emu(h))
    pic.crop_left = c.x
    pic.crop_top = c.y
    pic.crop_right = max(0.0, 1 - c.x - c.w)
    pic.crop_bottom = max(0.0, 1 - c.y - c.h)
    pic.name = b.block_id


def render_slide(prs, spec: SlideSpec, resolve: AssetResolver):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = next((b for b in spec.blocks if isinstance(b, ShapeBlock) and b.block_id == "bg"), None)
    if bg:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg.fill or "black")
    for b in spec.blocks:
        if isinstance(b, ShapeBlock):
            if b.block_id != "bg":
                _shape(slide, b)
        elif isinstance(b, TextBlock):
            _text(slide, b)
        elif isinstance(b, BarChartBlock):
            _bar_chart(slide, b)
        elif isinstance(b, LineChartBlock):
            _line_chart(slide, b)
        elif isinstance(b, TableBlock):
            _table(slide, b)
        elif isinstance(b, ImageBlock):
            _image(slide, b, resolve)
    script = spec.speaker_notes.strip()
    trace = spec.notes.strip()
    slide.notes_slide.notes_text_frame.text = (script + "\n\n— 근거 기록 —\n" + trace if script and trace else script or trace)
    return slide


def render_deck(deck: DeckSpec, out_path: Path, resolve: Optional[AssetResolver] = None) -> Path:
    prs = Presentation()
    canvas = tokens()["canvas"]
    prs.slide_width = _emu(canvas["w"])
    prs.slide_height = _emu(canvas["h"])
    prs.core_properties.title = deck.project_name
    prs.core_properties.subject = f"deck {deck.deck_id} v{deck.deck_version}"
    prs.core_properties.keywords = f"template {deck.template_version}; fingerprint {deck.input_fingerprint[:16]}"

    def _no_assets(asset_id: str):
        raise FileNotFoundError(f"asset {asset_id} not resolvable")

    for spec in deck.slides:
        render_slide(prs, spec, resolve or _no_assets)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
