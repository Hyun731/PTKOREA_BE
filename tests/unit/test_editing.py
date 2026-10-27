"""Deck editor commands: ids kept, evidence rules held, scope respected, Mock planner produces real edits."""
from datetime import date

import pytest

from core.composition.build import BuildInputs, build_deck
from core.editing import mock
from core.editing.commands import apply_commands
from core.editing.scope import AreaIn, SelectionIn, resolve_target
from core.editing.text import DeckVocabulary, TextEditError, from_tokens, to_tokens
from core.quality.qc import run_qc
from core.schemas.slidespec import ShapeBlock, TextBlock
from tests.helpers import load_fixture


@pytest.fixture(scope="module")
def built():
    p, f, po = load_fixture()
    return build_deck(BuildInputs(project=p, facts=f.facts + po.facts, posts=po.posts, generated_on=date(2026, 10, 12)))


def _fact_block(slides):
    for s in slides:
        for b in s.blocks:
            if isinstance(b, TextBlock) and any(p.type == "fact" for para in b.paragraphs for p in para.parts):
                return s, b
    raise AssertionError("fixture has no text block with a fact chip")


def test_tokens_round_trip_keeps_chips(built):
    slides = built.deck.slides
    s, b = _fact_block(slides)
    vocab = DeckVocabulary(slides)
    paras = from_tokens(to_tokens(b), b, vocab)
    assert [p.model_dump() for p in paras] == [p.model_dump() for p in b.paragraphs] or \
        ["".join(x.rendered() for x in p.parts) for p in paras] == ["".join(x.rendered() for x in p.parts) for p in b.paragraphs]


def test_text_edit_rejects_bare_numbers_and_unknown_facts(built):
    slides = built.deck.slides
    s, b = _fact_block(slides)
    vocab = DeckVocabulary(slides)
    with pytest.raises(TextEditError):
        from_tokens(["조회 3배 증가"], b, vocab)
    with pytest.raises(TextEditError):
        from_tokens(["{{F-" + "0" * 20 + "}} 증가"], b, vocab)
    res = apply_commands(slides, [{"op": "update_text", "slide_id": s.slide_id, "element_id": b.block_id, "paragraphs": ["12% 성장"]}])
    assert not res.applied and "근거 없는 숫자" in res.rejected[0]["reason"]
    assert res.slides[slides.index(s)].block(b.block_id).paragraphs == b.paragraphs  # untouched


def test_commands_keep_ids_and_only_touch_named_fields(built):
    slides = built.deck.slides
    s, b = _fact_block(slides)
    cmds = [
        {"op": "update_style", "slide_id": s.slide_id, "element_id": b.block_id, "size": b.size * 1.5, "color": "#1D4ED8", "reason": "강조"},
        {"op": "move_element", "slide_id": s.slide_id, "element_id": b.block_id, "x": 5000, "y": -20},
        {"op": "update_background", "slide_id": s.slide_id, "color": "tint"},
        {"op": "add_element", "slide_id": s.slide_id, "element": {"kind": "text", "x": 100, "y": 100, "w": 300, "h": 40, "paragraphs": ["메모"]}},
        {"op": "delete_element", "slide_id": s.slide_id, "element_id": "nope"},
    ]
    res = apply_commands(slides, cmds)
    assert [a["op"] for a in res.applied] == ["update_style", "move_element", "update_background", "add_element"]
    assert len(res.rejected) == 1
    out = next(x for x in res.slides if x.slide_id == s.slide_id)
    nb = out.block(b.block_id)
    assert nb.size == round(b.size * 1.5, 1) and nb.color == "#1D4ED8"
    assert nb.region.x + nb.region.w <= 1280 and nb.region.y == 0  # clamped to the canvas
    assert isinstance(out.blocks[0], ShapeBlock) and out.blocks[0].block_id == "bg"
    assert [p.fact_id for para in nb.paragraphs for p in para.parts if p.type == "fact"] == \
           [p.fact_id for para in b.paragraphs for p in para.parts if p.type == "fact"]
    others = [x for x in res.slides if x.slide_id != s.slide_id]
    assert others == [x for x in slides if x.slide_id != s.slide_id]  # nothing else changed
    qc = run_qc(built.deck.model_copy(update={"slides": res.slides}), built.metrics.registry, lambda a: built.assets.get(a))
    assert qc.blocks == run_qc(built.deck, built.metrics.registry, lambda a: built.assets.get(a)).blocks


def test_scope_selection_limits_targets(built):
    slides = built.deck.slides
    s, b = _fact_block(slides)
    r = b.region
    area = AreaIn(id="a", slide_id=s.slide_id, x=r.x / 1280, y=r.y / 720, w=r.w / 1280, h=r.h / 720)
    t = resolve_target(slides, s.slide_id, "selection", SelectionIn(areas=[area]))
    assert t.slide_ids == [s.slide_id] and b.block_id in t.element_ids[s.slide_id]
    other = next(x for x in s.blocks if x.block_id != b.block_id and x.block_id not in t.element_ids[s.slide_id])
    assert not t.allows(s.slide_id, other.block_id)
    assert resolve_target(slides, s.slide_id, "deck", SelectionIn()).slide_ids == [x.slide_id for x in slides]


def test_mock_planner_edits_the_canvas(built):
    slides = built.deck.slides
    s = next(x for x in slides if any(isinstance(b, TextBlock) and b.role == "lead" for b in x.blocks))
    t = resolve_target(slides, s.slide_id, "slide", SelectionIn())
    reply, cmds = mock.plan("제목 글씨를 키우고 파란색으로 바꿔줘", slides, t, s.slide_id)
    res = apply_commands(slides, cmds)
    lead = next(b for b in res.slides[slides.index(s)].blocks if getattr(b, "role", None) == "lead")
    before = next(b for b in s.blocks if getattr(b, "role", None) == "lead")
    assert res.applied and lead.color == "#1D4ED8" and lead.size > before.size and "파란색" in reply
    reply, cmds = mock.plan("무슨 말인지 모를 요청", slides, t, s.slide_id)
    assert cmds == [] and "Mock" in reply


def test_speaker_notes_and_rename(built, tmp_path):
    from pptx import Presentation
    from core.rendering.pptx_renderer import render_deck

    slides = built.deck.slides
    s = next(x for x in slides if x.page_role == "content")
    t = resolve_target(slides, s.slide_id, "slide", SelectionIn())
    reply, cmds = mock.plan("발표 대본 써줘", slides, t, s.slide_id)
    cmds.append({"op": "rename_slide", "slide_id": s.slide_id, "title": "새 이름"})
    res = apply_commands(slides, cmds)
    out = res.slides[slides.index(s)]
    assert out.speaker_notes.startswith("이번 장은") and out.title == "새 이름"
    assert out.blocks == s.blocks  # nothing on the slide changed
    path = render_deck(built.deck.model_copy(update={"slides": res.slides}), tmp_path / "d.pptx", lambda a: (built.assets[a].path, 1, 1))
    notes = Presentation(str(path)).slides[slides.index(s)].notes_slide.notes_text_frame.text
    assert notes.startswith("이번 장은") and "근거 기록" in notes
