"""Quantitative deck regression: order, split, QC 0 block, PPTX == SlideSpec (PRD 15, A09, A18)."""
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from core.cli import main as cli_main
from core.composition.build import BuildInputs, build_deck
from core.quality.qc import run_qc
from core.rendering.pptx_renderer import render_deck
from core.schemas.slidespec import BarChartBlock, LineChartBlock, TableBlock, TextBlock
from tests.helpers import FIXTURES, load_fixture


def build(posts_file="posts.csv", **kw):
    p, f, po = load_fixture(posts_file)
    return build_deck(BuildInputs(project=p, facts=f.facts + po.facts, posts=po.posts, generated_on=date(2026, 10, 12), **kw))


def order_key(slide_id: str) -> str:
    return "2-8c" if slide_id.startswith("2-8c-") else slide_id


def test_deck_order(expected):
    r = build()
    assert [order_key(s.slide_id) for s in r.deck.slides] == expected["quantitative_deck_order"]


def test_split_fixture(expected):
    r = build("posts_13.csv")
    creator = [s for s in r.deck.slides if s.slide_id.startswith("3-1-")]
    assert len(creator) == expected["split_fixture"]["creator_table_pages"]
    assert len(r.deck.slides) == expected["split_fixture"]["quantitative_total_pages"]
    assert creator[1].parent_slide_id == creator[0].slide_id
    rows = [len(b.rows) for s in creator for b in s.blocks if isinstance(b, TableBlock)]
    assert sum(rows) == 13 and max(rows) <= 12


@pytest.mark.parametrize("posts_file", ["posts.csv", "posts_13.csv"])
def test_qc_zero_blocks(expected, posts_file):
    r = build(posts_file)
    qc = run_qc(r.deck, r.metrics.registry, lambda a: None)
    assert qc.blocks == expected["acceptance"]["qc_blocks"], [i for i in qc.issues if i.severity == "block"]


def test_each_finding_leads_once_and_on_matching_page():
    r = build()
    leads = [fid for s in r.deck.slides if s.page_role == "content" for fid in s.finding_ids]
    assert sorted(leads) == sorted(f.id for f in r.findings)
    page_of = {fid: s.slide_id for s in r.deck.slides if s.page_role == "content" for fid in s.finding_ids}
    assert page_of["FN-R04-B1-경쟁X"] == "2-8a"
    assert page_of["FN-R03-B1-경쟁X"] == "2-8b"
    assert page_of["FN-R05-B1-경쟁Y-캐리어"] == "2-8c-캐리어"
    assert page_of["FN-R06-B1"] == "2-8d"
    assert page_of["FN-R10-B1"] == "3-1-B1"


def test_omitting_a_slot_removes_page_and_keeps_record():
    r = build(omitted_slots={"2-8d": "타겟 데이터 미확보로 생략"})
    assert "2-8d" not in [s.slide_id for s in r.deck.slides]
    slot = next(s for s in r.deck.plan if s.slot_id == "2-8d")
    assert slot.readiness == "omitted" and slot.reason


def _shape_text(shape) -> str:
    return "\n".join("".join(r.text for r in p.runs) for p in shape.text_frame.paragraphs)


def test_pptx_matches_slidespec(tmp_path: Path):
    """Same deck_version: every text/table/chart in the PPTX equals the SlideSpec (A18)."""
    r = build()
    path = render_deck(r.deck, tmp_path / "report.pptx")
    prs = Presentation(str(path))
    assert len(prs.slides) == len(r.deck.slides)
    for spec, slide in zip(r.deck.slides, prs.slides):
        shapes = {s.name: s for s in slide.shapes}
        for b in spec.blocks:
            if isinstance(b, TextBlock):
                expected = "\n".join(((p.bullet + " ") if p.bullet else "") + "".join(x.rendered() for x in p.parts)
                                     for p in b.paragraphs)
                assert _shape_text(shapes[b.block_id]) == expected, (spec.slide_id, b.block_id)
            elif isinstance(b, BarChartBlock):
                chart = shapes[b.block_id].chart  # native, editable chart
                assert list(chart.plots[0].categories) == [it.label for it in b.items]
                assert list(chart.series[0].values) == [it.value for it in b.items]
            elif isinstance(b, LineChartBlock):
                chart = shapes[b.block_id].chart
                assert [s.name for s in chart.series] == [s.entity for s in b.series]
            elif isinstance(b, TableBlock):
                table = shapes[b.block_id].table
                assert [c.text for c in table.rows[0].cells] == [c.label for c in b.columns]
                for r_i, row in enumerate(b.rows, start=1):
                    for c_i, cell in enumerate(row.cells):
                        if cell.merged:
                            continue
                        want = "\n".join("".join(x.rendered() for x in p.parts) for p in cell.paragraphs)
                        assert table.cell(r_i, c_i).text == want, (spec.slide_id, r_i, c_i)
        assert slide.notes_slide.notes_text_frame.text == spec.notes


def test_cli_exit_codes(tmp_path: Path):
    f = FIXTURES
    base = ["build", "--project", str(f / "project.json"), "--facts", str(f / "facts.csv"), "--posts", str(f / "posts.csv")]
    assert cli_main(base + ["--out", str(tmp_path / "ok"), "--accept-batches", "--date", "2026-10-12"]) == 0
    assert (tmp_path / "ok" / "report.pptx").exists()
    for name in ["factbook.csv", "sources.csv", "assets.csv", "findings.json", "deck_spec.json", "build_report.json", "errors.csv"]:
        assert (tmp_path / "ok" / name).exists(), name
    # Without batch acceptance nothing is eligible -> QC block.
    assert cli_main(base + ["--out", str(tmp_path / "blocked")]) == 2
    assert not (tmp_path / "blocked" / "report.pptx").exists()
    bad = tmp_path / "bad.json"
    bad.write_text("{}", encoding="utf-8")
    assert cli_main(["build", "--project", str(bad), "--facts", str(f / "facts.csv"), "--out", str(tmp_path / "x")]) == 1
