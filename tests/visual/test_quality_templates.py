"""Stage 0B structure checks for the 8 representative templates (QUALITY_BASELINE 2).

The human side-by-side comparison is produced by tools/quality_deck.py; these tests keep
the structural requirements from regressing.
"""
import os
import shutil
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from core.composition.build import BuildInputs, build_deck
from core.quality.qc import run_qc
from core.rendering.pptx_renderer import render_deck
from core.schemas.slidespec import BarChartBlock, ImageBlock, LineChartBlock, TableBlock, TextBlock
from tests.quality_fixture import build_inputs


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    out = tmp_path_factory.mktemp("quality")
    project, facts, posts, materials = build_inputs(out / "assets")
    r = build_deck(BuildInputs(project=project, facts=facts, posts=posts, materials=materials,
                               omitted_slots={"3-1-B2": "검증 범위 외", "3-1-B3": "검증 범위 외"},
                               generated_on=date(2026, 10, 12)))
    return r, out


def slide(r, slide_id):
    return next(s for s in r.deck.slides if s.slide_id == slide_id)


def test_qc_zero_blocks(result):
    r, _ = result
    qc = run_qc(r.deck, r.metrics.registry, lambda a: r.assets.get(a))
    assert qc.blocks == 0, [i for i in qc.issues if i.severity == "block"]


def test_p26_three_column_search_share(result):
    s = slide(result[0], "2-8a")
    charts = [b for b in s.blocks if isinstance(b, BarChartBlock)]
    assert len(charts) == 3
    for c in charts:
        values = [it.value for it in c.items]
        assert values == sorted(values, reverse=True)  # sorted bars
        assert sum(1 for it in c.items if it.role == "self") == 1
        assert sum(1 for it in c.items if it.role == "key") <= 1  # red only for the key rival


def test_p27_trend_panels_up_to_four_series_with_gap(result):
    s = slide(result[0], "2-8b")
    charts = [b for b in s.blocks if isinstance(b, LineChartBlock)]
    assert len(charts) == 3 and all(len(c.series) == 4 for c in charts)
    gaps = [p for c in charts for se in c.series for p in se.points if p.value is None]
    assert gaps, "missing month must stay a gap"


def test_p29_grouped_target_table(result):
    t = next(b for b in slide(result[0], "2-8d").blocks if isinstance(b, TableBlock))
    first_col = [row.cells[0] for row in t.rows]
    assert any(c.merge_down > 0 for c in first_col)  # brand group rows
    assert sum(1 for row in t.rows if row.cells[3].fill == "highlight") == 3  # brand-name rows


def test_p40_image_plus_table(result):
    s = slide(result[0], "3-1-B1")
    assert any(isinstance(b, ImageBlock) for b in s.blocks)
    t = next(b for b in s.blocks if isinstance(b, TableBlock))
    assert len(t.rows) == 11 and any(r.muted for r in t.rows)


def test_p54_p68_case_cards(result):
    r, _ = result
    # The fixture's cases are competitors': their card pages are the 경쟁사 분석 chapter's C1c.
    assert slide(r, "C1c").template_id == "case_card_4col.v1"
    assert slide(r, "C1c-p2").template_id == "case_card_3col.v1"
    for sid in ("C1c", "C1c-p2"):
        s = slide(r, sid)
        frames = [b for b in s.blocks if getattr(b, "block_id", "").endswith("_frame")]
        assert len({round(f.region.h) for f in frames}) == 1  # aligned card heights
        assert all(a in r.assets for a in s.asset_ids)


def test_p89_recap_three_pillars_no_new_facts(result):
    r, _ = result
    s = slide(r, "5-recap")
    heads = [b for b in s.blocks if isinstance(b, TextBlock) and b.block_id.endswith("_head")]
    assert len(heads) == 3
    finding_facts = {fid for f in r.findings for fid in f.fact_ids}
    assert {b.fact_id for b in s.bindings} <= finding_facts


def test_p104_source_table_links(result, tmp_path):
    r, out = result
    s = slide(r, "A-source")
    t = next(b for b in s.blocks if isinstance(b, TableBlock))
    links = [p for row in t.rows for p in row.cells[2].paragraphs[0].parts if p.type == "link"]
    assert links and all(p.href.startswith("https://") for p in links)
    path = render_deck(r.deck, tmp_path / "v.pptx", lambda a: (r.assets[a].path, r.assets[a].width, r.assets[a].height))
    prs = Presentation(str(path))
    idx = [x.slide_id for x in r.deck.slides].index("A-source")
    table = next(sh for sh in prs.slides[idx].shapes if sh.name == "table").table
    run = table.cell(1, 2).text_frame.paragraphs[0].runs[0]
    assert run.hyperlink.address.startswith("https://")  # clickable in PowerPoint


@pytest.mark.skipif(not shutil.which("docker") or os.environ.get("SKIP_RENDER") == "1", reason="renderer not available")
def test_real_render_page_count(result, tmp_path):
    from core.rendering.preview import render_pages
    r, _ = result
    path = render_deck(r.deck, tmp_path / "v.pptx", lambda a: (r.assets[a].path, r.assets[a].width, r.assets[a].height))
    pages = render_pages(path, tmp_path / "preview")
    assert len(pages) == len(r.deck.slides)
