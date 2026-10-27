"""Stage 1 flow through the API: manual materials -> PPT, with review, stale and QC gates
(PRD A05, A06, A11 partial, A13, A15, A16, A17; DEVELOPMENT_PLAN 단계 1 완료 조건)."""
import io
import os

from PIL import Image
from pptx import Presentation

from tests.integration.conftest import HAS_RENDER


def png_bytes(w=1600, h=900) -> bytes:
    img = Image.new("RGB", (w, h), (40, 90, 160))
    for x in range(0, w, 100):
        for y in range(h // 2, h):
            img.putpixel((x, y), (255, 255, 0))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def ok(r, status=200):
    assert r.status_code == status, r.text
    return r.json()


def test_full_manual_flow(client):
    p = ok(client.post("/v1/projects/sample"), 201)
    pid = p["id"]
    plan = ok(client.get(f"/v1/projects/{pid}/research-plan"))
    sections = {(i["section_id"], i["role"]) for i in plan["items"]}
    assert ("search", "self") in sections and ("campaign", "competitor") in sections
    assert all(c["status"] in ("key_missing", "manual_only") for i in plan["items"] for c in i["channels"])

    # Automatic research is honest about missing keys (no fake run).
    campaign_item = next(i for i in plan["items"] if i["section_id"] == "campaign")
    r = client.post(f"/v1/projects/{pid}/research-runs", json={"item_ids": [campaign_item["id"]]})
    assert r.status_code == 503 and r.json()["error"]["code"] == "PROVIDER_NOT_CONFIGURED"

    # CSV batches are candidates until accepted.
    batches = ok(client.post(f"/v1/projects/{pid}/imports/sample"), 201)
    assert all(b["status"] == "candidate" and b["synthetic"] for b in batches[:1])
    sp = ok(client.get(f"/v1/projects/{pid}/slide-plan"))
    assert sp["counts"]["generated"] == 0
    for b in batches:
        ok(client.post(f"/v1/imports/{b['id']}/review", json={"action": "accept"}))
    sp = ok(client.get(f"/v1/projects/{pid}/slide-plan"))
    ready = {s["slot_id"] for s in sp["plan"] if s["readiness"] == "ready"}
    assert {"2-8a", "2-8b", "2-8d", "3-1-B1", "5-recap"} <= ready
    c1 = next(s for s in sp["plan"] if s["slot_id"] == "C1")
    assert c1["readiness"] == "needs_input"

    # Manual case material: URL + comment + image + crop + metric fact + claims.
    item = next(i for i in plan["items"] if i["section_id"] == "campaign" and i["entity"] == "경쟁X")
    m = ok(client.post(f"/v1/projects/{pid}/materials", json={
        "title": "경쟁X 팝업 캠페인", "url": "https://example.com/campaign/x", "publisher": "예시 매체",
        "published_at": "2026-03-02", "brand_id": "B1", "entity": "경쟁X", "section_id": "campaign",
        "label": "경쟁X · 팝업 연계", "comment": "이 표현 방식을 참고", "activity_date": "2026-03-02",
        "research_item_id": item["id"]}), 201)
    bad = client.post(f"/v1/projects/{pid}/materials", json={"title": "x", "url": "file:///etc/passwd"})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "INVALID_URL"

    a = ok(client.post(f"/v1/materials/{m['id']}/assets", files={"file": ("case.png", png_bytes(), "image/png")},
                       data={"title": "팝업 현장"}), 201)
    nonimg = client.post(f"/v1/materials/{m['id']}/assets", files={"file": ("x.png", b"not an image", "image/png")})
    assert nonimg.status_code == 415
    a = ok(client.patch(f"/v1/assets/{a['id']}/crop", json={"base_version": a["version"], "crop": {"x": 0.1, "y": 0.0, "w": 0.8, "h": 1.0}}))
    assert a["crop"]["w"] == 0.8 and len(a["crop_history"]) == 1
    conflict = client.patch(f"/v1/assets/{a['id']}/crop", json={"base_version": 1, "crop": {"x": 0, "y": 0, "w": 1, "h": 1}})
    assert conflict.status_code == 409 and conflict.json()["error"]["code"] == "STALE_VERSION"
    ok(client.patch(f"/v1/assets/{a['id']}", json={"review_status": "accepted"}))

    f = ok(client.post(f"/v1/materials/{m['id']}/facts", json={"metric": "views", "label": "조회", "value": "116799",
                                                               "measured_at": "2026-09-30"}), 201)
    ok(client.patch(f"/v1/facts/{f['id']}/review", params={"project_id": pid}, json={"action": "confirm"}))

    c_fact = ok(client.post(f"/v1/materials/{m['id']}/claims", json={
        "kind": "fact", "text": "신제품 공개 후 같은 비주얼로 [[센텀 팝업2]] 운영"}), 201)
    assert c_fact["issues"] == [] and any(p["type"] == "name" for p in c_fact["parts"])
    c_bad = ok(client.post(f"/v1/materials/{m['id']}/claims", json={"kind": "fact", "text": "조회 200만 달성"}), 201)
    assert c_bad["issues"]
    r = client.patch(f"/v1/claims/{c_bad['id']}", json={"base_version": c_bad["version"], "review_status": "accepted"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "UNBOUND_NUMBER"
    ok(client.patch(f"/v1/claims/{c_bad['id']}", json={"base_version": c_bad["version"], "review_status": "excluded"}))
    c_int = ok(client.post(f"/v1/materials/{m['id']}/claims", json={
        "kind": "interpretation", "text": "조회 {{" + f["id"] + "}} 규모의 체험형 확산"}), 201)
    assert c_int["fact_ids"] == [f["id"]] and "116,799" in c_int["text"]
    for c in (c_fact, c_int):
        ok(client.patch(f"/v1/claims/{c['id']}", json={"base_version": c["version"], "review_status": "accepted"}))
    r = client.patch(f"/v1/materials/{m['id']}/review", json={"status": "excluded"})
    assert r.status_code == 422  # exclusion needs a reason
    m = ok(client.patch(f"/v1/materials/{m['id']}/review", json={"status": "accepted"}))
    assert m["readiness"] == []

    sp = ok(client.get(f"/v1/projects/{pid}/slide-plan"))
    assert next(s for s in sp["plan"] if s["slot_id"] == "C1")["readiness"] == "ready"
    blocking = [i for i in sp["qc"]["issues"] if i["severity"] == "block"]
    # M1 has no accepted market evidence: the person omits it explicitly.
    for i in blocking:
        assert i["code"] == "SLOT_NEEDS_INPUT", i
        ok(client.put(f"/v1/projects/{pid}/slot-omissions/{i['block_id']}", json={"reason": "이번 보고서 범위 외"}))
    sp = ok(client.get(f"/v1/projects/{pid}/slide-plan"))
    assert sp["qc"]["blocks"] == 0, sp["qc"]

    # Deck build (job) -> state recoverable from the DB.
    start = ok(client.post(f"/v1/projects/{pid}/deck-builds", headers={"Idempotency-Key": "build-1"}), 202)
    job = ok(client.get(f"/v1/jobs/{start['job']['id']}"))
    assert job["state"] == "succeeded", job
    again = ok(client.post(f"/v1/projects/{pid}/deck-builds", headers={"Idempotency-Key": "build-1"}), 202)
    assert again["reused"] and again["deck_id"] == start["deck_id"]
    deck = ok(client.get(f"/v1/decks/{start['deck_id']}"))
    ids = [s["slide_id"] for s in deck["spec"]["slides"]]
    assert "C1" in ids and "A-source" in ids and deck["qc"]["blocks"] == 0 and not deck["stale"]
    c1 = next(s for s in deck["spec"]["slides"] if s["slide_id"] == "C1")
    assert c1["asset_ids"] == [a["id"]] and m["source"]["id"] in c1["source_ids"]

    if HAS_RENDER:
        assert deck["render_status"] == "ok" and deck["pages"] == len(ids)
        assert client.get(f"/v1/decks/{deck['id']}/preview/1").headers["content-type"] == "image/png"
        exp_job = ok(client.post(f"/v1/decks/{deck['id']}/exports", headers={"Idempotency-Key": "exp-1"}), 202)["job"]
        assert exp_job["state"] == "succeeded", exp_job
        exports = ok(client.get(f"/v1/projects/{pid}/exports"))
        assert {"report.pptx", "factbook.csv", "sources.csv", "assets.csv", "build_report.json"} <= set(exports[0]["files"])
        pptx = client.get(f"/v1/exports/{exports[0]['id']}/files/report.pptx")
        prs = Presentation(io.BytesIO(pptx.content))
        slide = prs.slides[ids.index("C1")]
        pic = next(s for s in slide.shapes if s.name == "card0_image")
        assert abs(pic.crop_left - 0.1) < 1e-6 and abs(pic.crop_right - 0.1) < 1e-6  # same crop as the editor
        texts = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        assert "116,799" in texts and "센텀 팝업2" in texts
        sources_csv = client.get(f"/v1/exports/{exports[0]['id']}/files/sources.csv").text
        assert "https://example.com/campaign/x" in sources_csv
    else:
        r = client.post(f"/v1/decks/{deck['id']}/exports")
        assert r.status_code == 409 and r.json()["error"]["code"] == "RENDER_FAILED"

    # Evidence changes -> dependent claim stale, deck stale, export blocked (A16).
    ok(client.patch(f"/v1/facts/{f['id']}/review", params={"project_id": pid},
                    json={"action": "correct", "value": "116800", "note": "원문 재확인"}))
    mat = ok(client.get(f"/v1/materials/{m['id']}"))
    assert next(c for c in mat["claims"] if c["id"] == c_int["id"])["review_status"] == "stale"
    assert "근거 변경 재검토" in mat["readiness"]
    deck = ok(client.get(f"/v1/decks/{start['deck_id']}"))
    assert deck["stale"]
    r = client.post(f"/v1/decks/{deck['id']}/exports")
    assert r.status_code == 409 and r.json()["error"]["code"] == "STALE_VERSION"

    # Optimistic concurrency on project settings (A11).
    proj = ok(client.get(f"/v1/projects/{pid}"))
    ok(client.patch(f"/v1/projects/{pid}", json={"base_version": proj["version"], "changes": {"target_hypothesis": "35+ 가족 여행"}}))
    r = client.patch(f"/v1/projects/{pid}", json={"base_version": proj["version"], "changes": {"target_hypothesis": "x"}})
    assert r.status_code == 409


def test_project_validation(client):
    r = client.post("/v1/projects", json={"name": "x", "client": "y", "period": {"start": "2026-01-01", "end": "2026-09-30"},
                                          "brands": [{"id": "b1", "name": "A", "competitors": ["B"]}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_PROJECT"
    r = client.post("/v1/projects", json={
        "name": "새 프로젝트", "client": "고객", "period": {"start": "2026-01-01", "end": "2026-09-30"},
        "brands": [{"id": "B1", "name": "자사", "competitors": ["경쟁1", "경쟁2"], "categories": ["가방"],
                    "target": {"age_min": 20, "age_max": 34, "gender": "female", "label": "여성 20~34세"}}]})
    body = ok(r, 201)
    assert body["comparison_window"]["start"] == "2025-10-01"
    assert {w["section_id"] for w in body["config"]["section_windows"]} >= {"influencer", "campaign", "market"}
    # No search-index provider: search slides are opt-in, never planned by default.
    assert not any(t.startswith("2-8") for t in body["config"]["enabled_slides"])
    assert "search" not in {i["section_id"] for i in ok(client.get(f"/v1/projects/{body['id']}/research-plan"))["items"]}
    meta = ok(client.get("/v1/meta"))
    assert [p["id"] for p in meta["profiles"]] == ["diagnostic"]  # proposal not exposed in P0


def test_posts_batch_replaces_only_same_brand(client):
    p = ok(client.post("/v1/projects", json={
        "name": "브랜드별 배치", "client": "고객", "period": {"start": "2026-01-01", "end": "2026-09-30"},
        "brands": [{"id": "B1", "name": "자사", "competitors": ["경쟁1", "경쟁2"], "categories": ["가방"]},
                   {"id": "B2", "name": "자사2", "competitors": ["경쟁3", "경쟁4"], "categories": ["가방"]}]}), 201)
    pid = p["id"]
    head = "post_id,brand,channel,creator,owner_type,posted_at,measured_at,followers,followers_measured_at,views,likes,comments,likes_hidden,product,url\n"

    def upload(brand, post):
        row = f"{post},{brand},youtube,크리에이터 {post},creator,2026-06-05,2026-10-01,500000,2026-10-01,1000000,200,50,false,가방,\n"
        b = ok(client.post(f"/v1/projects/{pid}/imports", data={"kind": "posts"},
                           files={"file": (f"{post}.csv", (head + row).encode(), "text/csv")}), 201)
        return ok(client.post(f"/v1/imports/{b['id']}/review", json={"action": "accept"}))["id"]

    b1, b2, b1_again = upload("B1", "P1"), upload("B2", "P2"), upload("B1", "P3")
    status = {b["id"]: b["status"] for b in ok(client.get(f"/v1/projects/{pid}/imports"))}
    assert status == {b1: "replaced", b2: "accepted", b1_again: "accepted"}


def test_deck_editor_copy_ai_and_export(client):
    pid = ok(client.post("/v1/projects/sample"), 201)["id"]
    for b in ok(client.post(f"/v1/projects/{pid}/imports/sample"), 201):
        ok(client.post(f"/v1/imports/{b['id']}/review", json={"action": "accept"}))
    deck_id = ok(client.post(f"/v1/projects/{pid}/deck-builds"), 202)["deck_id"]
    ed = ok(client.get(f"/v1/decks/{deck_id}/edit"))
    assert not ed["edited"] and ed["slides"] and "mock" in ed["ai"]["modes"]
    slides = ed["slides"]
    sid = next(s["slide_id"] for s in slides if s["page_role"] == "content" and any(b.get("role") == "lead" for b in s["blocks"]))

    # Mock AI: commands are previewed (slides returned), nothing is saved until the editor saves.
    out = ok(client.post(f"/v1/decks/{deck_id}/edit/ai", json={
        "message": "제목 글씨를 키우고 파란색으로 바꿔줘", "slides": slides, "current_slide_id": sid, "scope": "slide", "mode": "mock"}))
    assert out["mode"] == "mock" and out["applied"] and sid in out["changed"]
    assert not ok(client.get(f"/v1/decks/{deck_id}/edit"))["edited"]
    # A selection scope keeps the AI inside the marked element.
    sel = ok(client.post(f"/v1/decks/{deck_id}/edit/ai", json={
        "message": "파란색으로 바꿔줘", "slides": slides, "current_slide_id": sid, "scope": "selection",
        "selection": {"element_ids": {sid: ["lead"]}}, "mode": "mock"}))
    assert {a["element_id"] for a in sel["applied"]} == {"lead"}

    # Save the previewed slides; optimistic version; QC still holds numbers to evidence.
    merged = [next((x for x in out["slides"] if x["slide_id"] == s["slide_id"]), s) for s in slides]
    saved = ok(client.put(f"/v1/decks/{deck_id}/edit", json={"base_version": ed["edit_version"], "slides": merged, "title": "편집본"}))
    assert client.put(f"/v1/decks/{deck_id}/edit", json={"base_version": ed["edit_version"], "slides": merged}).status_code == 409
    bad = [dict(s) for s in merged]
    lead = next(b for b in next(s for s in bad if s["slide_id"] == sid)["blocks"] if b.get("role") == "lead")
    lead["paragraphs"] = [{"parts": [{"type": "text", "text": "조회 3배"}]}]
    qc = ok(client.post(f"/v1/decks/{deck_id}/edit/qc", json={"slides": bad}))
    assert any(i["code"] == "UNBOUND_NUMBER" and i["slide_id"] == sid for i in qc["issues"])

    exp = ok(client.post(f"/v1/decks/{deck_id}/edit/export"), 201) if not saved["qc"]["blocks"] else None
    if exp:
        assert "report.pptx" in exp["files"]
        r = client.get(f"/v1/exports/{exp['id']}/files/report.pptx")
        assert r.status_code == 200 and len(r.content) > 10_000
    reset = ok(client.delete(f"/v1/decks/{deck_id}/edit"))
    assert not reset["edited"] and reset["edit_version"] == saved["edit_version"] + 1


def test_editor_ai_log_apply_and_restore(client):
    pid = ok(client.post("/v1/projects/sample"), 201)["id"]
    for b in ok(client.post(f"/v1/projects/{pid}/imports/sample"), 201):
        ok(client.post(f"/v1/imports/{b['id']}/review", json={"action": "accept"}))
    deck_id = ok(client.post(f"/v1/projects/{pid}/deck-builds"), 202)["deck_id"]
    ed = ok(client.get(f"/v1/decks/{deck_id}/edit"))
    slides = ed["slides"]
    sid = next(s["slide_id"] for s in slides if s["page_role"] == "content" and any(b.get("role") == "lead" for b in s["blocks"]))

    def lead(sl):
        s = next(x for x in sl if x["slide_id"] == sid)
        return next(b for b in s["blocks"] if b.get("role") == "lead")

    out = ok(client.post(f"/v1/decks/{deck_id}/edit/ai", json={"message": "제목 파란색으로 바꿔줘", "slides": slides,
                                                                "current_slide_id": sid, "scope": "slide", "mode": "mock"}))
    log = ok(client.get(f"/v1/decks/{deck_id}/edit/log"))
    assert log[0]["id"] == out["log_id"] and log[0]["status"] == "proposed" and log[0]["mode"] == "mock"
    merged = [next((x for x in out["slides"] if x["slide_id"] == s["slide_id"]), s) for s in slides]
    applied = ok(client.post(f"/v1/decks/{deck_id}/edit/log/{out['log_id']}/apply",
                             json={"base_version": ed["edit_version"], "slides_before": slides, "slides": merged}))
    assert applied["log"]["status"] == "applied" and applied["log"]["can_restore_before"]
    assert client.post(f"/v1/decks/{deck_id}/edit/log/{out['log_id']}/apply",
                       json={"base_version": applied["edit_version"], "slides_before": slides, "slides": merged}).status_code == 409
    assert lead(ok(client.get(f"/v1/decks/{deck_id}/edit"))["slides"])["color"] == "#1D4ED8"

    # a cancelled proposal stays in the log but changes nothing
    out2 = ok(client.post(f"/v1/decks/{deck_id}/edit/ai", json={"message": "제목 키워줘", "slides": merged,
                                                                 "current_slide_id": sid, "scope": "slide", "mode": "mock"}))
    ok(client.post(f"/v1/decks/{deck_id}/edit/log/{out2['log_id']}/cancel"))

    # go back to before the applied edit, then undo that restore
    back = ok(client.post(f"/v1/decks/{deck_id}/edit/log/{out['log_id']}/restore", json={"base_version": applied["edit_version"], "to": "before"}))
    assert lead(back["slides"])["color"] == lead(slides)["color"] and back["log"]["kind"] == "restore"
    again = ok(client.post(f"/v1/decks/{deck_id}/edit/log/{back['log']['id']}/restore", json={"base_version": back["edit_version"], "to": "before"}))
    assert lead(again["slides"])["color"] == "#1D4ED8"
    kinds = [(e["kind"], e["status"]) for e in ok(client.get(f"/v1/decks/{deck_id}/edit/log"))]
    assert kinds == [("restore", "applied"), ("restore", "applied"), ("ai", "cancelled"), ("ai", "applied")]


def test_backend_token_required_when_set(client, monkeypatch):
    from dataclasses import replace

    from app import main

    monkeypatch.setattr(main, "settings", replace(main.settings, backend_token="s3cret"))
    r = client.get("/v1/meta")
    assert r.status_code == 401 and r.json()["error"]["code"] == "UNAUTHORIZED"
    assert client.get("/v1/meta", headers={"X-Backend-Token": "wrong"}).status_code == 401
    assert client.get("/openapi.json").status_code == 401
    assert client.get("/v1/meta", headers={"X-Backend-Token": "s3cret"}).status_code == 200
    assert client.get("/healthz").status_code == 200
